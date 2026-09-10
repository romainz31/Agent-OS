from __future__ import annotations

import json
import os
import re
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any


class DocumentPipelineError(RuntimeError):
    pass


class DocumentPipelineBuilder:
    """Build and verify a real XLSX from indexed document rows."""

    MONEY_RE = re.compile(
        r"(?<!\d)(\d{1,3}(?:[ .]\d{3})*(?:[,.]\d{1,2})|\d+[,.]\d{1,2})\s*(€|EUR)?",
        re.IGNORECASE,
    )

    DATE_PATTERNS = (
        re.compile(r"\b(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})\b"),
        re.compile(r"\b(20\d{2})-(\d{2})-(\d{2})\b"),
    )

    def __init__(self, workspace: str | Path) -> None:
        self.workspace = Path(workspace).resolve()

    @staticmethod
    def _clean(value: Any) -> str:
        return " ".join(str(value or "").strip().split())

    @staticmethod
    def _safe_json(value: Any) -> dict[str, Any]:
        if isinstance(value, dict):
            return dict(value)
        if not isinstance(value, str) or not value.strip():
            return {}
        try:
            parsed = json.loads(value)
        except Exception:
            return {}
        return parsed if isinstance(parsed, dict) else {}

    @staticmethod
    def _safe_float(value: Any) -> float | None:
        if value in (None, ""):
            return None
        if isinstance(value, (int, float)):
            return float(value)
        text = str(value).strip().replace("\u202f", " ").replace(" ", "")
        text = text.replace("€", "").replace("EUR", "").replace("eur", "")
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        else:
            text = text.replace(",", ".")
        try:
            return float(text)
        except Exception:
            return None

    @classmethod
    def _extract_dates(cls, text: str) -> list[str]:
        result: list[str] = []
        raw = str(text or "")
        for index, pattern in enumerate(cls.DATE_PATTERNS):
            for match in pattern.finditer(raw):
                try:
                    if index == 0:
                        day = int(match.group(1))
                        month = int(match.group(2))
                        year = int(match.group(3))
                        if year < 100:
                            year += 2000
                    else:
                        year = int(match.group(1))
                        month = int(match.group(2))
                        day = int(match.group(3))
                    value = datetime(year, month, day).date().isoformat()
                except Exception:
                    continue
                if value not in result:
                    result.append(value)
                if len(result) >= 12:
                    return result
        return result

    @classmethod
    def _extract_amounts(cls, text: str) -> list[float]:
        result: list[float] = []
        for match in cls.MONEY_RE.finditer(str(text or "")):
            value = cls._safe_float(match.group(1))
            if value is None:
                continue
            if value not in result:
                result.append(value)
            if len(result) >= 30:
                break
        return result

    @staticmethod
    def _flatten(value: Any, prefix: str = "", out=None):
        if out is None:
            out = []
        if isinstance(value, dict):
            for key, child in value.items():
                path = f"{prefix}.{key}" if prefix else str(key)
                if isinstance(child, (dict, list)):
                    DocumentPipelineBuilder._flatten(child, path, out)
                else:
                    out.append((path, str(child if child is not None else "")))
            return out
        if isinstance(value, list):
            for index, child in enumerate(value, 1):
                path = f"{prefix}[{index}]" if prefix else f"[{index}]"
                if isinstance(child, (dict, list)):
                    DocumentPipelineBuilder._flatten(child, path, out)
                else:
                    out.append((path, str(child if child is not None else "")))
            return out
        out.append((prefix or "value", str(value if value is not None else "")))
        return out

    @staticmethod
    def _iso_date(value: Any) -> str:
        text = str(value or "").strip()
        if not text:
            return ""
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y"):
            try:
                return datetime.strptime(text[:10], fmt).date().isoformat()
            except Exception:
                pass
        return text[:32]

    @classmethod
    def _invoice(cls, analysis: dict[str, Any]) -> dict[str, Any]:
        invoice = analysis.get("invoice")
        if not isinstance(invoice, dict):
            invoice = {}

        ht = cls._safe_float(
            invoice.get("subtotal_ex_tax")
            or invoice.get("total_ht")
            or invoice.get("ht")
        )
        tax = cls._safe_float(invoice.get("tax") or invoice.get("tva"))
        ttc = cls._safe_float(
            invoice.get("total_including_tax")
            or invoice.get("total_ttc")
            or invoice.get("ttc")
        )

        vat_rate = None
        if ht not in (None, 0) and tax is not None:
            vat_rate = round((tax / ht) * 100, 2)

        issue = cls._iso_date(invoice.get("issue_date") or invoice.get("date"))
        due = cls._iso_date(invoice.get("due_date") or invoice.get("echeance"))

        delay = None
        try:
            if issue and due:
                delay = (
                    datetime.fromisoformat(due[:10]).date()
                    - datetime.fromisoformat(issue[:10]).date()
                ).days
        except Exception:
            delay = None

        return {
            "supplier": cls._clean(
                invoice.get("supplier")
                or invoice.get("vendor")
                or invoice.get("fournisseur")
            ),
            "customer": cls._clean(invoice.get("customer") or invoice.get("client")),
            "invoice_number": cls._clean(
                invoice.get("invoice_number")
                or invoice.get("number")
                or invoice.get("numero")
            ),
            "issue_date": issue,
            "due_date": due,
            "subtotal_ex_tax": ht,
            "tax": tax,
            "total_including_tax": ttc,
            "currency": cls._clean(
                invoice.get("currency")
                or ("EUR" if any(v is not None for v in (ht, tax, ttc)) else "")
            ),
            "vat_rate_pct": vat_rate,
            "payment_delay_days": delay,
        }

    @staticmethod
    def _list_text(value: Any) -> str:
        if isinstance(value, list):
            values = []
            for item in value:
                if isinstance(item, dict):
                    item = item.get("value") or item.get("name") or json.dumps(item, ensure_ascii=False)
                clean = " ".join(str(item or "").split()).strip()
                if clean and clean not in values:
                    values.append(clean)
            return " ; ".join(values)
        if isinstance(value, dict):
            return json.dumps(value, ensure_ascii=False)
        return " ".join(str(value or "").split()).strip()

    def prepare(self, rows: list[dict[str, Any]]) -> dict[str, Any]:
        records = []
        details = []

        for row in rows:
            analysis = self._safe_json(row.get("analysis_json"))
            extracted = str(row.get("extracted") or "")
            invoice = self._invoice(analysis)

            dates = []
            structured_dates = analysis.get("dates")
            if isinstance(structured_dates, list):
                for item in structured_dates:
                    if isinstance(item, dict):
                        item = item.get("value") or item.get("date")
                    value = self._iso_date(item)
                    if value and value not in dates:
                        dates.append(value)
            for value in self._extract_dates(extracted):
                if value not in dates:
                    dates.append(value)

            amounts = []
            structured_amounts = analysis.get("amounts")
            if isinstance(structured_amounts, list):
                for item in structured_amounts:
                    if isinstance(item, dict):
                        item = item.get("value") or item.get("amount") or item.get("montant")
                    number = self._safe_float(item)
                    if number is not None and number not in amounts:
                        amounts.append(number)
            for number in self._extract_amounts(extracted):
                if number not in amounts:
                    amounts.append(number)

            source = str(row.get("path") or "")
            path = Path(source)
            try:
                size = int(path.stat().st_size) if path.is_file() else None
            except Exception:
                size = None

            image = analysis.get("image") if isinstance(analysis.get("image"), dict) else {}
            local_image = (
                analysis.get("local_image_features")
                if isinstance(analysis.get("local_image_features"), dict)
                else {}
            )

            record = {
                "file_name": path.name,
                "source": source,
                "extension": path.suffix.lower(),
                "file_size_bytes": size,
                "category": self._clean(analysis.get("category") or row.get("category") or "autres"),
                "document_type": self._clean(analysis.get("document_type")),
                "summary": self._clean(analysis.get("summary") or row.get("summary")),
                **invoice,
                "dates_found": " ; ".join(dates[:12]),
                "amounts_found": " ; ".join(f"{v:.2f}" for v in amounts[:30]),
                "amount_count": len(amounts),
                "min_amount": min(amounts) if amounts else None,
                "max_amount": max(amounts) if amounts else None,
                "dominant_color": self._clean(
                    image.get("dominant_color_family") or local_image.get("dominant_family")
                ),
                "dominant_hex": self._clean(
                    image.get("dominant_hex") or local_image.get("dominant_hex")
                ),
                "tags": self._list_text(analysis.get("tags")),
                "extraction_chars": len(extracted),
                "partial_extract": bool(row.get("limited")),
                "structured_analysis": bool(analysis),
            }
            records.append(record)

            flat = self._flatten(analysis)
            if not flat:
                flat = [
                    ("summary", record["summary"]),
                    ("category", record["category"]),
                ]
            for key, value in flat:
                details.append({"source": source, "field": key, "value": value})

        category_counts: dict[str, int] = {}
        ttc_values = []
        for record in records:
            category = record["category"] or "autres"
            category_counts[category] = category_counts.get(category, 0) + 1
            if record["total_including_tax"] is not None:
                ttc_values.append(float(record["total_including_tax"]))

        return {
            "records": records,
            "details": details,
            "metrics": {
                "document_count": len(records),
                "structured_count": sum(1 for r in records if r["structured_analysis"]),
                "category_counts": category_counts,
                "invoice_with_ttc_count": len(ttc_values),
                "invoice_total_ttc_sum": round(sum(ttc_values), 2) if ttc_values else None,
            },
        }

    def output_path(self, mission_ref: str, requested: str | None = None) -> Path:
        # Ne crée aucun dossier pendant la phase de demande d'autorisation.
        # Le répertoire de sortie n'est créé que pendant create_xlsx(), après accord.
        export_dir = (self.workspace / "documents_exports").resolve()

        if requested:
            requested_path = Path(requested)
            if requested_path.is_absolute():
                resolved = requested_path.resolve()
                try:
                    resolved.relative_to(self.workspace)
                except ValueError as exc:
                    raise DocumentPipelineError(
                        "Pour cette version, le XLSX de sortie doit rester dans le workspace."
                    ) from exc
            else:
                name = requested_path.name
                if not name.lower().endswith(".xlsx"):
                    name += ".xlsx"
                resolved = export_dir / name
        else:
            clean_ref = re.sub(r"[^A-Za-z0-9_-]+", "_", str(mission_ref or "mission"))
            resolved = export_dir / f"synthese_{clean_ref}.xlsx"

        base = resolved
        index = 2
        while resolved.exists():
            resolved = base.with_name(f"{base.stem}_{index}{base.suffix}")
            index += 1
        return resolved

    def create_xlsx(
        self,
        output: str | Path,
        prepared: dict[str, Any],
        *,
        objective: str,
        mission_ref: str,
    ) -> dict[str, Any]:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError as exc:
            raise DocumentPipelineError(
                "openpyxl absent : installe requirements-documents.txt."
            ) from exc

        output = Path(output).resolve()
        try:
            output.relative_to(self.workspace)
        except ValueError as exc:
            raise DocumentPipelineError("Destination hors workspace refusée.") from exc

        if output.exists():
            raise DocumentPipelineError(f"Le fichier existe déjà : {output}")

        output.parent.mkdir(parents=True, exist_ok=True)

        book = Workbook()
        ws = book.active
        ws.title = "Synthese"

        columns = [
            ("file_name", "Fichier"),
            ("source", "Source"),
            ("extension", "Extension"),
            ("file_size_bytes", "Taille octets"),
            ("category", "Catégorie"),
            ("document_type", "Type"),
            ("supplier", "Fournisseur"),
            ("customer", "Client"),
            ("invoice_number", "N° facture"),
            ("issue_date", "Date émission"),
            ("due_date", "Échéance"),
            ("subtotal_ex_tax", "HT"),
            ("tax", "TVA"),
            ("total_including_tax", "TTC"),
            ("currency", "Devise"),
            ("vat_rate_pct", "Taux TVA calculé %"),
            ("payment_delay_days", "Délai paiement calculé j"),
            ("dates_found", "Dates détectées"),
            ("amounts_found", "Montants détectés"),
            ("amount_count", "Nb montants"),
            ("min_amount", "Montant min"),
            ("max_amount", "Montant max"),
            ("dominant_color", "Couleur dominante"),
            ("dominant_hex", "Couleur HEX"),
            ("tags", "Tags"),
            ("summary", "Synthèse"),
            ("extraction_chars", "Caractères extraits"),
            ("partial_extract", "Extrait partiel"),
            ("structured_analysis", "Analyse structurée"),
        ]

        ws.append([label for _, label in columns])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for record in prepared.get("records", []):
            ws.append([record.get(key, "") for key, _ in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        details = book.create_sheet("Details")
        details.append(["Source", "Champ", "Valeur"])
        for cell in details[1]:
            cell.font = Font(bold=True)
        for item in prepared.get("details", []):
            details.append([item.get("source", ""), item.get("field", ""), item.get("value", "")])

        mission = book.create_sheet("Mission")
        mission.append(["Champ", "Valeur"])
        for cell in mission[1]:
            cell.font = Font(bold=True)
        metrics = prepared.get("metrics", {})
        for key, value in (
            ("Mission", mission_ref),
            ("Objectif", objective),
            ("Documents", metrics.get("document_count")),
            ("Analyses structurées", metrics.get("structured_count")),
            ("Factures avec TTC", metrics.get("invoice_with_ttc_count")),
            ("Somme TTC structurée", metrics.get("invoice_total_ttc_sum")),
            ("Catégories", json.dumps(metrics.get("category_counts", {}), ensure_ascii=False)),
        ):
            mission.append([key, value])

        for sheet in book.worksheets:
            for cells in sheet.columns:
                letter = cells[0].column_letter
                width = max(
                    [10]
                    + [
                        min(len(str(cell.value or "")) + 2, 72)
                        for cell in cells[:250]
                    ]
                )
                sheet.column_dimensions[letter].width = width

        fd, temp_name = tempfile.mkstemp(
            prefix=".agentos-xlsx-",
            suffix=".tmp",
            dir=output.parent,
        )
        os.close(fd)
        temp = Path(temp_name)
        try:
            # Document text is data, including text beginning with '='.
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        if cell.data_type == 'f':
                            cell.data_type = 's'
            book.save(temp)
            if output.exists():
                raise DocumentPipelineError(f"Le fichier existe déjà : {output}")
            os.replace(temp, output)
        finally:
            temp.unlink(missing_ok=True)

        verification = self.inspect_xlsx(output)
        if not verification.get("valid"):
            output.unlink(missing_ok=True)
            raise DocumentPipelineError(
                "XLSX invalide : " + str(verification.get("error") or "vérification échouée")
            )

        return {
            "path": str(output),
            "metrics": metrics,
            "verification": verification,
        }

    def inspect_xlsx(self, path: str | Path) -> dict[str, Any]:
        path = Path(path)
        if not path.is_file():
            return {"valid": False, "error": "Fichier introuvable."}

        try:
            with zipfile.ZipFile(path, "r") as archive:
                if archive.testzip() is not None:
                    return {"valid": False, "error": "Archive XLSX corrompue."}
        except Exception as exc:
            return {"valid": False, "error": str(exc)}

        try:
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=False)
            sheets = list(book.sheetnames)
            rows_per_sheet = {}
            non_empty_cells = 0
            for sheet in book.worksheets:
                non_empty_rows = 0
                for row in sheet.iter_rows():
                    has_value = False
                    for cell in row:
                        if cell.value not in (None, ""):
                            has_value = True
                            non_empty_cells += 1
                    if has_value:
                        non_empty_rows += 1
                rows_per_sheet[sheet.title] = non_empty_rows
            book.close()
            return {
                "valid": bool(sheets) and non_empty_cells > 0,
                "sheets": sheets,
                "rows_per_sheet": rows_per_sheet,
                "non_empty_cells": non_empty_cells,
                "size_bytes": int(path.stat().st_size),
                "error": "",
            }
        except Exception as exc:
            return {"valid": False, "error": str(exc)}
