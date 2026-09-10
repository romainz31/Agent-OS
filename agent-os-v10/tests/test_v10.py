"""Offline behavioural checks. No personal data or live model needed."""
import io
import tempfile
import unittest
from pathlib import Path
from datetime import datetime,date
from zoneinfo import ZoneInfo
from unittest.mock import patch
from agentos.agenda import PersonalAgenda
from agentos.life_journal import LifeJournal
from agentos.file_assistant import FileAssistant

class Checks(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.now=datetime(2026,9,10,9,0,tzinfo=ZoneInfo('Europe/Paris'))
        self.a=PersonalAgenda(Path(self.tmp.name)/'agenda.db',now_provider=lambda:self.now)
        self.j=LifeJournal(self.a)
    def tearDown(self): self.tmp.cleanup()
    def test_empty(self): self.assertEqual(self.j.rows('1970-01-01','9999-12-31'),[])
    def test_multiple_and_dedup(self):
        text="Journal : nettoyé le filtre; arrosé les plantes; fait les courses"
        self.j.handle(text);self.j.handle(text)
        self.assertEqual(len(self.j.rows('2026-09-10','2026-09-10')),3)
    def test_restart(self):
        self.j.add('filtre','2026-05-01')
        self.assertEqual(len(LifeJournal(self.a).rows('2026-01-01','2026-09-10')),1)
    def test_natural_jai_list_is_stored_and_recalled(self):
        answer=self.j.handle("aujourdhui jai nettoyé la fontaine a chats, les toilettes, la machine a café")
        self.assertIn("tu as nettoyé la fontaine a chats",answer)
        rows=self.j.completed_rows('2026-09-10','2026-09-10')
        self.assertEqual([row['title'] for row in rows],[
            'nettoyé la fontaine a chats',
            'nettoyé les toilettes',
            'nettoyé la machine a café',
        ])
        recall=self.j.handle("j'ai fait quoi comme taches aujourdhui?")
        self.assertIn('tu as fait 3 choses',recall.lower())
        self.assertIn('nettoyé la machine a café',recall)

    def test_timeline_uses_agenda_items_as_single_store(self):
        self.j.handle("jai nettoyé la fontaine a chats")
        self.a._add_item(kind='todo',title='nettoyer la piscine',due_date=date(2026,9,14),start_time=None,source_text='lundi je dois nettoyer la piscine')
        self.a._add_item(kind='appointment',title='dentiste',due_date=date(2026,9,14),start_time='15:00',source_text='rdv dentiste lundi à 15h')
        with self.a._connect() as db:
            kinds={row['kind'] for row in db.execute('SELECT kind FROM agenda_items').fetchall()}
            legacy=db.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='life_events'").fetchone()
        self.assertTrue({'action','todo','appointment'} <= kinds)
        self.assertIsNone(legacy)
        monday=self.j.timeline_rows('2026-09-14','2026-09-14')
        self.assertEqual({row['kind'] for row in monday},{'todo','appointment'})

    def test_legacy_life_events_are_migrated_once(self):
        with tempfile.TemporaryDirectory() as t:
            agenda=PersonalAgenda(Path(t)/'agenda.db',now_provider=lambda:self.now)
            with agenda._connect() as db:
                db.execute("CREATE TABLE life_events(id INTEGER PRIMARY KEY, day TEXT NOT NULL, title TEXT NOT NULL, kind TEXT NOT NULL, source TEXT NOT NULL, created TEXT NOT NULL)")
                db.execute('INSERT INTO life_events VALUES(1,?,?,?,?,?)',(
                    '2026-09-09','nettoyé le filtre','action','ancien journal',self.now.isoformat()))
            timeline=LifeJournal(agenda)
            self.assertEqual(len(timeline.completed_rows('2026-09-09','2026-09-09')),1)
            LifeJournal(agenda)
            with agenda._connect() as db:
                count=db.execute("SELECT COUNT(*) FROM agenda_items WHERE kind='action'").fetchone()[0]
            self.assertEqual(count,1)

    def test_non_action_jai_is_not_logged(self):
        self.assertIsNone(self.j.handle("jai faim"))
        self.assertEqual(self.j.completed_rows('2026-09-10','2026-09-10'),[])

    def test_count_months(self):
        self.j.add('nettoyé le filtre','2026-06-09')
        self.j.add('nettoyé le filtre','2026-06-10')
        self.j.add('nettoyé le filtre','2026-08-10')
        answer=self.j.handle("Combien de fois j'ai nettoyé le filtre sur les trois derniers mois ?")
        self.assertIn('2 fois',answer)
        self.assertIn('tu as nettoyé le filtre',answer)
    def test_future_negation_question(self):
        self.j.handle("J'ai pas nettoyé le filtre")
        self.j.handle("J'ai nettoyé le filtre ?")
        self.j.handle('Journal : demain nettoyer le filtre')
        self.assertEqual(self.j.rows('1970-01-01','9999-12-31'),[])
    def test_rollover_and_appointments(self):
        for kind in ('todo','appointment'):
            self.a._add_item(kind=kind,title='test '+kind,due_date=date(2026,9,9),start_time='10:00',source_text='le 09/09/2026')
        self.j.rollover();self.j.rollover()
        with self.a._connect() as db:
            todo=db.execute("SELECT * FROM agenda_items WHERE kind='todo'").fetchone()
            ap=db.execute("SELECT * FROM agenda_items WHERE kind='appointment'").fetchone()
        self.assertEqual(todo['due_date'],'2026-09-10');self.assertEqual(ap['due_date'],'2026-09-09')
        self.assertEqual(self.a._todo_metadata(dict(todo))['reschedule_count'],1)
    def test_undated_not_moved(self):
        self.a._add_item(kind='todo',title='test',due_date=date(2026,9,9),start_time=None,source_text='À faire : test')
        self.j.rollover()
        self.assertEqual(self.a.backlog_todos()[0]['due_date'],'2026-09-09')
    def test_completion_single_source(self):
        self.a._add_item(kind='todo',title='nettoyer le filtre',due_date=date(2026,9,10),start_time=None,source_text="aujourd'hui nettoyer le filtre")
        self.j.handle("J'ai nettoyé le filtre")
        self.assertEqual(len(self.a.backlog_todos()),0)
        self.assertEqual(len(self.j.rows('2026-09-10','2026-09-10')),1)
        self.j.handle("J'ai nettoyé le filtre")
        self.assertEqual(len(self.j.rows('2026-09-10','2026-09-10')),1)
    def test_mood_and_delete(self):
        answer=self.j.handle('Humeur : fatigué mais content')
        self.assertIn("Je vois : fatigué mais content",answer)
        self.assertEqual(self.j.rows('1970-01-01','9999-12-31')[0]['kind'],'humeur')
        self.j.handle('journal supprimer 1')
        self.assertEqual(self.j.rows('1970-01-01','9999-12-31'),[])
    def test_notices_once(self):
        self.assertEqual(len(self.j.notifications()),1)
        self.assertEqual(self.j.notifications(),[])
    def test_pdf_scan_and_text(self):
        from reportlab.pdfgen import canvas
        stream=io.BytesIO(); c=canvas.Canvas(stream)
        c.drawString(60,720,'Invoice number A123. Total amount 120 euros.');c.showPage();c.showPage();c.save()
        service=FileAssistant(None,Path(self.tmp.name)/'work',Path(self.tmp.name)/'docs')
        service.vision=lambda raw,objective:'SCAN total TTC 240 EUR'
        try:
            text,limited=service.extract(stream.getvalue(),'.pdf')
            self.assertIn('[Page 1]',text);self.assertIn('[Page 2]',text);self.assertIn('240 EUR',text);self.assertFalse(limited)
        finally: service.close()
    def test_xlsx_read(self):
        from openpyxl import Workbook
        b=Workbook();b.active.append(['Facture',120]);out=io.BytesIO();b.save(out)
        service=FileAssistant(None,Path(self.tmp.name)/'work',Path(self.tmp.name)/'docs')
        try:self.assertIn('Facture | 120',service.extract(out.getvalue(),'.xlsx')[0])
        finally:service.close()

    def test_xlsx_export_literals(self):
        from agentos.document_pipeline import DocumentPipelineBuilder
        from openpyxl import load_workbook
        builder=DocumentPipelineBuilder(self.tmp.name)
        out=Path(self.tmp.name)/'export.xlsx'
        builder.create_xlsx(out,{'records':[{'supplier':'=1+1','total_including_tax':120}], 'details':[], 'metrics':{}},objective='test',mission_ref='M-001')
        book=load_workbook(out)
        self.assertEqual(book['Synthese']['G2'].value,'=1+1')
        self.assertEqual(book['Synthese']['G2'].data_type,'s')
        self.assertEqual(book['Synthese']['N2'].value,120)
        book.close()

    def test_telegram_attachment_authorization(self):
        from telegram_client import TelegramAgentOS
        from unittest.mock import Mock
        bot=TelegramAgentOS(token='test',agentos_url='http://127.0.0.1:8765',allowed_chat_id=42)
        bot.send=Mock();bot.receive_attachment=Mock()
        bot.handle_message({'chat':{'id':99},'document':{'file_id':'secret'}})
        bot.receive_attachment.assert_not_called()
        bot.handle_message({'chat':{'id':42},'photo':[{'file_id':'small'},{'file_id':'large'}]})
        self.assertEqual(bot.receive_attachment.call_args.args[1]['file_id'],'large')

    def test_pending_attachment_survives_restart(self):
        path=Path(self.tmp.name)/'work'/'photo.jpg'
        path.parent.mkdir()
        path.write_bytes(b'image')
        service=FileAssistant(None,Path(self.tmp.name)/'work',Path(self.tmp.name)/'docs')
        service.hold_attachment('telegram-42',path,'mon-chat.jpg','image')
        service.close()
        reopened=FileAssistant(None,Path(self.tmp.name)/'work',Path(self.tmp.name)/'docs')
        try:
            pending=reopened.pending_attachment('telegram-42')
            self.assertEqual(pending['original_name'],'mon-chat.jpg')
            self.assertTrue(reopened.clear_pending_attachment('telegram-42'))
            self.assertIsNone(reopened.pending_attachment('telegram-42'))
        finally:
            reopened.close()

if __name__=='__main__':unittest.main()
